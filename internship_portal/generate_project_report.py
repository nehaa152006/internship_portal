"""
generate_project_report.py
--------------------------
Generates a comprehensive, highly professional Microsoft Word (.docx) document
for the University Internship Management & Tracking System (T&P Portal).
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # 1. Page Margins (1 inch everywhere)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette
    COLOR_PRIMARY = RGBColor(0, 102, 255)     # #0066FF Royal Blue
    COLOR_DARK = RGBColor(15, 23, 42)          # #0F172A Dark Slate
    COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B Slate Muted
    COLOR_SUCCESS = RGBColor(22, 101, 52)      # #166534 Emerald
    HEX_HEADER_BG = "0066FF"
    HEX_LIGHT_ROW = "F8FAFC"
    HEX_BORDER = "E2E8F0"

    # Default Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(51, 65, 85)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # COVER / TITLE SECTION
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run("UNIVERSITY TRAINING & PLACEMENT SECTION")
    title_run.font.size = Pt(13)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    main_heading = doc.add_paragraph()
    main_heading.paragraph_format.space_after = Pt(12)
    h_run = main_heading.add_run("Internship Management & Tracking System\n(T&P Portal)")
    h_run.font.size = Pt(26)
    h_run.font.bold = True
    h_run.font.color.rgb = COLOR_DARK
    main_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(30)
    sub_run = subtitle_p.add_run("A Complete University Placement, Verification, Supervisor Feedback & Campus Opportunity Platform")
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = COLOR_MUTED
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Project Details Table Box on Title Page
    info_table = doc.add_table(rows=6, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    details = [
        ("Project Title:", "University Internship Management & Tracking System"),
        ("Architecture:", "Python / Flask MVC Architecture with SQLite3 Database"),
        ("Key Features:", "Letter Verification, Milestone Tracking, Supervisor Scorecards, Campus Drives Push, Candidate Shortlisting & Interview Scheduling (When & Where)"),
        ("User Roles:", "1. Student (Candidate)  |  2. Training & Placement Officer (Admin)"),
        ("UI/UX Design:", "Minimalist University Interface, Google Font Inter, Bootstrap 5"),
        ("Document Type:", "Technical Software Documentation & Academic Project Report")
    ]

    for idx, (label, val) in enumerate(details):
        row = info_table.rows[idx]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_margins(c0, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_DARK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)

    doc.add_page_break()

    # Helper function for Section Headings
    def add_chapter_heading(number_str, title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(20)
        h.paragraph_format.space_after = Pt(8)
        run_num = h.add_run(f"{number_str}. ")
        run_num.font.bold = True
        run_num.font.size = Pt(16)
        run_num.font.color.rgb = COLOR_PRIMARY

        run_title = h.add_run(title_str)
        run_title.font.bold = True
        run_title.font.size = Pt(16)
        run_title.font.color.rgb = COLOR_DARK
        return h

    def add_subheading(title_str):
        sh = doc.add_paragraph()
        sh.paragraph_format.space_before = Pt(14)
        sh.paragraph_format.space_after = Pt(4)
        r = sh.add_run(title_str)
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return sh

    def add_bullet(p, bold_prefix, text_content):
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(bold_prefix + ": ")
        r1.font.bold = True
        r1.font.color.rgb = COLOR_DARK
        r2 = bp.add_run(text_content)
        return bp

    # -------------------------------------------------------------
    # CHAPTER 1: EXECUTIVE SUMMARY & ABSTRACT
    # -------------------------------------------------------------
    add_chapter_heading("1", "Executive Summary & Project Overview")
    
    p = doc.add_paragraph()
    p.add_run(
        "The University Internship Management & Tracking System is an enterprise-grade web application tailored for the "
        "Training & Placement (T&P) Section of academic institutions. Internships represent a mandatory, high-stakes milestone "
        "in modern higher education, requiring strict documentation, verification of authentic employer credentials, continuous "
        "mentor evaluation, and official academic credit clearance. Historically, academic placement departments have operated using "
        "disparate spreadsheets, paper physical records, and uncoordinated email chains, leading to communication bottlenecks, lost "
        "documents, and zero transparency for students."
    )
    
    p2 = doc.add_paragraph()
    p2.add_run(
        "This project resolves these critical bottlenecks by introducing a centralized, dual-panel digital portal. Students are equipped "
        "with an intuitive, self-service dashboard to register company details, upload mandatory verification documents (Offer Letter, "
        "Completion Certificate, Project Report), track milestone progression in real time, view supervisor feedback, download official "
        "credit clearance NOC certificates, and apply for campus recruitment drives. Concurrently, the T&P Administrator is provided with "
        "a command center to audit submissions, request corrections, record supervisor competency scorecards, broadcast recruitment drives, "
        "and schedule interviews with explicit When (Date & Time) and Where (Venue / Link) specifications visible directly to students."
    )

    # -------------------------------------------------------------
    # CHAPTER 2: PROBLEM STATEMENT & OBJECTIVES
    # -------------------------------------------------------------
    add_chapter_heading("2", "Problem Statement, Vision & Objectives")
    
    add_subheading("2.1 Current System Challenges")
    add_bullet(doc, "Data Fragmentation", "Student internship records, letters, and emails are scattered across personal inboxes and unstructured spreadsheets.")
    add_bullet(doc, "Lack of Transparency", "Students frequently visit placement offices in person to inquire whether their offer letter has been approved or rejected.")
    add_bullet(doc, "Missing Industry Feedback", "Universities struggle to systematically collect and record qualitative competency ratings from company mentors.")
    add_bullet(doc, "Uncoordinated Recruitment Drives", "Announcing internship openings and tracking candidate shortlisting and interview venues lacks a single source of truth.")
    add_bullet(doc, "Delay in Credit Clearances", "Compiling completion certificates for final degree grade cards is delayed due to disorganized manual paperwork.")

    add_subheading("2.2 Proposed Solution & Strategic Objectives")
    add_bullet(doc, "Complete Digital Lifecycle", "Unify the entire student journey from registration to final academic NOC clearance in one transparent system.")
    add_bullet(doc, "Two-Way Transparency", "Ensure every action taken by the T&P officer (approvals, remarks, correction requests, interview invites) is immediately visible to the student.")
    add_bullet(doc, "Campus Opportunity Broadcast", "Allow T&P officers to push drives, receive one-click student applications with tailored resumes, and manage applicant pipelines.")
    add_bullet(doc, "Clear Interview Scheduling", "Ensure shortlisted candidates receive exact interview logistics (Date, Time, Location/Link, Instructions) on their dashboard and via notifications.")

    # -------------------------------------------------------------
    # CHAPTER 3: SYSTEM ARCHITECTURE & TECHNOLOGY STACK
    # -------------------------------------------------------------
    add_chapter_heading("3", "System Architecture & Technology Stack")

    p = doc.add_paragraph()
    p.add_run(
        "The application is engineered on the Model-View-Controller (MVC) architectural pattern, chosen for its modularity, "
        "simplicity of explanation in technical defenses, and high performance without the overhead of heavy third-party ORMs."
    )

    tech_table = doc.add_table(rows=7, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    t_headers = ["Layer / Component", "Technology Selected", "Role & Engineering Purpose"]
    for i, h in enumerate(t_headers):
        cell = tech_table.rows[0].cells[i]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    tech_data = [
        ("Web Framework", "Python 3 & Flask", "Handles routing, session authentication, HTTP requests, file uploads, and role security."),
        ("Database Engine", "SQLite3 (sqlite3 module)", "Lightweight, zero-configuration relational database with raw parameterized queries for transparent execution."),
        ("Templating Engine", "Jinja2 (bundled with Flask)", "Dynamic server-side HTML rendering, template inheritance, conditionals, and filters."),
        ("Security & Auth", "Werkzeug Security", "Cryptographic password hashing (generate_password_hash) and secure file upload sanitization (secure_filename)."),
        ("Front-End UI", "HTML5, CSS3, Bootstrap 5", "Responsive grid layout, modern cards, modal dialogs, and university-grade minimalist design system."),
        ("Design Assets", "FontAwesome 6 & Google Fonts", "Inter font typography and official vector icons matching modern placement dashboards.")
    ]

    for row_idx, data in enumerate(tech_data, start=1):
        row = tech_table.rows[row_idx]
        bg = HEX_LIGHT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.font.bold = True

    # -------------------------------------------------------------
    # CHAPTER 4: DATABASE DESIGN & SCHEMA
    # -------------------------------------------------------------
    add_chapter_heading("4", "Database Design & Data Dictionary")

    p = doc.add_paragraph()
    p.add_run(
        "The relational database schema is stored in 'internship_portal.db' and managed via 'database.py'. "
        "Foreign key constraints and strict CHECK constraints ensure referential integrity across the five primary tables:"
    )

    add_subheading("4.1 'users' Table (Authentication & Profile)")
    p = doc.add_paragraph("Stores credentials and student profiles (Roll Number, Department, CGPA, Phone, Master Resume).")
    
    # Users table
    utbl = doc.add_table(rows=7, cols=3)
    utbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    u_headers = ["Field Name", "Data Type & Constraints", "Description"]
    for i, h in enumerate(u_headers):
        c = utbl.rows[0].cells[i]
        set_cell_background(c, "1E293B")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    u_rows = [
        ("id", "INTEGER PRIMARY KEY AUTOINCREMENT", "Unique identifier for the user account."),
        ("full_name", "TEXT NOT NULL", "Candidate's legal academic name."),
        ("email", "TEXT UNIQUE NOT NULL", "University institutional email address used for login."),
        ("password_hash", "TEXT NOT NULL", "Werkzeug secure salted password hash."),
        ("role", "TEXT NOT NULL CHECK(role IN ('student', 'admin'))", "Access role determining dashboard permissions."),
        ("roll_number, branch, cgpa, phone", "TEXT", "Academic profile fields used during drive applications.")
    ]
    for r_idx, r_data in enumerate(u_rows, 1):
        for c_idx, val in enumerate(r_data):
            cell = utbl.rows[r_idx].cells[c_idx]
            set_cell_background(cell, HEX_LIGHT_ROW if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            cell.paragraphs[0].add_run(val).font.size = Pt(9)

    add_subheading("4.2 'internships' Table (Core Tracking & Evaluation)")
    p = doc.add_paragraph("Stores comprehensive company information, verification status, uploaded file paths, and mentor ratings.")
    
    add_bullet(doc, "Company Profile", "company_name, company_location, company_website, role, domain, mode (Offline/Online/Hybrid), internship_type (Paid/Unpaid), stipend.")
    add_bullet(doc, "Tenure & Contact", "start_date, end_date, manager_name, manager_email.")
    add_bullet(doc, "Document Tracking", "letter_filename (Offer Letter), completion_filename, report_filename with status checks ('Pending', 'Submitted', 'Uploaded').")
    add_bullet(doc, "Review & Evaluation", "review_status ('Pending', 'Under Review', 'Approved', 'Ongoing', 'Completed', 'Correction Required', 'Rejected'), admin_remarks.")
    add_bullet(doc, "Competency Scorecard", "feedback_rating (1-5 stars), tech_skills (1-5), communication (1-5), teamwork (1-5), problem_solving (1-5), professionalism (1-5), strengths, improvements.")

    add_subheading("4.3 'events' & 'event_applications' Tables (Pushed Drives & Interviews)")
    p = doc.add_paragraph(
        "Allows the T&P Officer to push recruitment drives to all students, and allows students to submit applications. "
        "Crucially, the 'event_applications' table records the scheduled interview details:"
    )
    add_bullet(doc, "status", "'Applied', 'Under Review', 'Shortlisted', 'Selected', 'Rejected'.")
    add_bullet(doc, "interview_date", "Date when the student's interview will take place (e.g., 2026-10-15).")
    add_bullet(doc, "interview_time", "Time of the interview (e.g., 11:00 AM).")
    add_bullet(doc, "interview_venue", "Physical room/location or online video link (Google Meet/MS Teams).")
    add_bullet(doc, "interview_instructions", "Preparation instructions (e.g., carry 2 copies of resume, college ID, and laptop).")
    add_bullet(doc, "admin_notes", "Officer feedback and remarks regarding student screening.")

    # -------------------------------------------------------------
    # CHAPTER 5: STUDENT PANEL MODULE
    # -------------------------------------------------------------
    doc.add_page_break()
    add_chapter_heading("5", "Student Panel: Architecture & User Experience")

    p = doc.add_paragraph()
    p.add_run(
        "The Student Panel has been crafted with a clean, university-grade minimalist design system. It is accessed via "
        "a persistent sidebar navigation matching institutional placement portals:"
    )

    add_subheading("5.1 Dashboard ('/student/dashboard')")
    p = doc.add_paragraph(
        "Serves as the central mission control for the student. Features include:"
    )
    add_bullet(doc, "Shortlisted Interview Banner", "If the candidate is shortlisted for any drive, a prominent callout banner immediately highlights When (Date & Time) and Where (Venue / Link) the interview is scheduled, accompanied by specific instructions.")
    add_bullet(doc, "Summary Metrics", "Four clean metric cards highlighting Verification Status, Internship State, Documents Submitted (X / 3), and Supervisor Evaluation Rating.")
    add_bullet(doc, "Active Internship Snapshot", "Displays company name, position, duration, and T&P correction warnings if any action is needed.")
    add_bullet(doc, "Visual 5-Stage Roadmap", "A milestone tracker illustrating Added -> Submitted -> Verified -> Ongoing -> Completed stages.")
    add_bullet(doc, "Pushed Opportunities", "Shows the top active campus drives with instant 'Apply Now' access.")

    add_subheading("5.2 My Internship ('/student/internship')")
    p = doc.add_paragraph(
        "Allows the student to review their complete company profile, mentor contact info, and offer letter status. "
        "If details are saved as a draft, or if the T&P officer flags 'Correction Required', the student can edit the details inline "
        "and resubmit for re-verification."
    )

    add_subheading("5.3 Document Center ('/student/documents')")
    p = doc.add_paragraph(
        "Dedicated upload and preview management center for the 3 mandatory university documents: "
        "1) Offer / Approval Letter, 2) Completion Certificate, and 3) Final Internship Project Report. "
        "Files can be previewed or replaced seamlessly."
    )

    add_subheading("5.4 Progress & Milestone Tracker ('/student/progress')")
    p = doc.add_paragraph(
        "An educational lifecycle breakdown answering three vital student questions at all times: "
        "1) What is completed?, 2) What is pending?, and 3) What action should you take next?"
    )

    add_subheading("5.5 Supervisor Feedback Scorecard ('/student/feedback')")
    p = doc.add_paragraph(
        "Displays the official industry supervisor evaluation, featuring an overall 5-star rating, qualitative review comments, "
        "and 5 competency progress bars (Technical Skills, Communication, Teamwork, Problem Solving, Professionalism)."
    )

    add_subheading("5.6 Final Evaluation & Credit Clearance ('/student/completion')")
    p = doc.add_paragraph(
        "Provides an academic clearance checklist and generates an official printable University NOC & Internship Clearance Certificate "
        "suitable for college credit transfer with window.print() integration."
    )

    add_subheading("5.7 Campus Drives & Events ('/student/events')")
    p = doc.add_paragraph(
        "Lists all drives published by T&P. Students can click 'Apply Now' to submit their application with custom pitch notes and tailored resumes. "
        "Once shortlisted, the card displays a dedicated 'Interview Scheduled' box detailing the date, time, venue, and instructions."
    )

    add_subheading("5.8 In-App Notifications ('/student/notifications')")
    p = doc.add_paragraph(
        "Real-time notifications keeping the candidate updated regarding offer approvals, admin remarks, new published drives, "
        "and interview scheduling details with a 'Mark All as Read' feature."
    )

    add_subheading("5.9 Student Profile & Master Resume ('/student/profile')")
    p = doc.add_paragraph(
        "Enables students to update their academic details (Roll No, Department, CGPA, Phone) and upload a master resume PDF that "
        "automatically populates drive applications."
    )

    # -------------------------------------------------------------
    # CHAPTER 6: T&P ADMIN PANEL MODULE
    # -------------------------------------------------------------
    add_chapter_heading("6", "T&P Administrator (Placement Officer) Panel")

    p = doc.add_paragraph()
    p.add_run(
        "The Administrator Panel empowers Training & Placement Officers to govern the student body, verify authentic industry "
        "letters, communicate requirements, push campus placement drives, and coordinate interviews:"
    )

    add_subheading("6.1 Tabbed Admin Dashboard ('/admin/dashboard')")
    add_bullet(doc, "Tab 1: Internship Verifications", "Displays aggregate metrics (Total Records, Letters Submitted, Letters Pending, Completions Uploaded), status filter buttons, search bar, and full student record table.")
    add_bullet(doc, "Tab 2: Pushed Events & Drives", "Displays all active recruitment opportunities, company names, deadlines, statuses (Open/Closed), and student applicant counts.")

    add_subheading("6.2 Verification & Evaluation ('/admin/internship/<id>')")
    p = doc.add_paragraph(
        "The officer can inspect uploaded letters, verify company details, and execute approval workflows: "
        "Approve, Reject, or set 'Correction Required' with custom admin remarks explaining what the student needs to fix. "
        "The officer can also enter the industry manager feedback rating and competency scores."
    )

    add_subheading("6.3 Publishing Events & Drives ('/admin/events/push')")
    p = doc.add_paragraph(
        "Form for the officer to publish internship drives and T&P requests. Fields include Title, Company Name, Event Type, Role, "
        "Eligibility, Stipend, Mode/Location, Deadline, Description, and Link. An 'Instant Student Broadcast' checkbox automatically "
        "delivers notifications to all student accounts."
    )

    add_subheading("6.4 Candidate Shortlisting & Interview Scheduling ('/admin/events/<id>')")
    p = doc.add_paragraph(
        "The officer can inspect all student applicants for any drive, download their resumes, read custom notes, and click "
        "'Update & Schedule' to select their status (Shortlisted, Selected, Rejected). When shortlisting, the officer enters the "
        "Interview Date, Time, Venue/Link, and Instructions, which triggers an automated notification and dashboard alert for the candidate."
    )

    # -------------------------------------------------------------
    # CHAPTER 7: SECURITY & ENGINEERING BEST PRACTICES
    # -------------------------------------------------------------
    add_chapter_heading("7", "Security, Reliability & Code Architecture")

    add_bullet(doc, "Password Security", "All passwords are encrypted using PBKDF2 with SHA-256 via Werkzeug Security. Plain text passwords are never stored.")
    add_bullet(doc, "Role-Based Access Control (RBAC)", "Protected routes enforce @login_required and @admin_required decorators to prevent unauthorized access.")
    add_bullet(doc, "Secure File Uploads", "File uploads are sanitized with secure_filename() and restricted to an allowed extension whitelist (PDF, JPG, PNG). Unique user ID prefixes prevent file collision.")
    add_bullet(doc, "SQL Injection Prevention", "All database transactions use parameterized queries with tuple parameters (?), neutralizing SQL injection vulnerabilities.")
    add_bullet(doc, "Session Management", "Flask secure signed cookies track authenticated session states with secret key signing.")

    # -------------------------------------------------------------
    # CHAPTER 8: EXECUTION & VERIFICATION GUIDE
    # -------------------------------------------------------------
    add_chapter_heading("8", "Installation, Execution & Demo Guide")

    p = doc.add_paragraph("To run the University Internship Management System locally:")
    
    code_box = doc.add_paragraph()
    code_box.paragraph_format.left_indent = Inches(0.5)
    code_box.paragraph_format.space_before = Pt(4)
    code_box.paragraph_format.space_after = Pt(10)
    c_run = code_box.add_run(
        "# 1. Navigate to project root\n"
        "cd internship_portal\n\n"
        "# 2. Install required dependencies\n"
        "pip install -r requirements.txt\n\n"
        "# 3. Run database initialization & migrations (if starting fresh)\n"
        "python migrate3.py\n\n"
        "# 4. Launch the Flask application server\n"
        "python app.py\n"
        "# Server will start at: http://127.0.0.1:5000"
    )
    c_run.font.name = 'Consolas'
    c_run.font.size = Pt(9.5)
    c_run.font.color.rgb = RGBColor(30, 41, 59)

    add_subheading("8.1 Seeded Demo Accounts")
    demo_table = doc.add_table(rows=3, cols=3)
    demo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    demo_table.rows[0].cells[0].paragraphs[0].add_run("User Role").font.bold = True
    demo_table.rows[0].cells[1].paragraphs[0].add_run("Email Address").font.bold = True
    demo_table.rows[0].cells[2].paragraphs[0].add_run("Password").font.bold = True
    for c in demo_table.rows[0].cells:
        set_cell_background(c, "1E293B")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    creds = [
        ("Student Account", "student@demo.com", "student123"),
        ("T&P Administrator", "admin@university.edu", "admin123")
    ]
    for idx, (role, em, pw) in enumerate(creds, 1):
        row = demo_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(role).font.bold = True
        row.cells[1].paragraphs[0].add_run(em)
        row.cells[2].paragraphs[0].add_run(pw)
        for c in row.cells:
            set_cell_background(c, HEX_LIGHT_ROW if idx == 1 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)

    # -------------------------------------------------------------
    # CHAPTER 9: VIVA VOCE & TECHNICAL DEFENSE Q&A
    # -------------------------------------------------------------
    add_chapter_heading("9", "Viva Voce & Technical Defense Q&A")

    qa_list = [
        ("Why did you select SQLite over MySQL or PostgreSQL for this project?",
         "SQLite requires zero external server setup, runs as a single serverless file, and perfectly demonstrates core SQL concepts (transactions, foreign keys, constraints) during live demos and viva examinations. The code is modular, meaning it can easily be swapped for PostgreSQL in a high-scale deployment."),
        
        ("How does the system ensure students know where and when their interview will take place?",
         "When an administrator updates an applicant's status to 'Shortlisted', the portal provides specialized fields for Interview Date, Time, Venue/Link, and Instructions. Upon submission, these fields are stored in the 'event_applications' table, immediately triggering an in-app notification and rendering an alert banner on both the Student Dashboard and Events page."),
        
        ("What happens if an offer letter has discrepancies?",
         "The administrator sets the review status to 'Correction Required' and enters specific remarks in the 'admin_remarks' field. On the student dashboard, a prominent warning alert highlights the remarks and offers an 'Edit & Resubmit' button that changes the status back to 'Pending' once updated."),
        
        ("How are file uploads securely isolated?",
         "Files are categorized into dedicated directories (uploads/letters, uploads/completion, uploads/reports, uploads/resumes). Each file is processed using secure_filename() and prepended with the student's unique ID to prevent namespace collisions and directory traversal exploits.")
    ]

    for q, a in qa_list:
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(8)
        qp.paragraph_format.space_after = Pt(2)
        r_q = qp.add_run(f"Q: {q}")
        r_q.font.bold = True
        r_q.font.color.rgb = COLOR_PRIMARY
        
        ap = doc.add_paragraph()
        ap.paragraph_format.space_after = Pt(6)
        r_a = ap.add_run(f"A: {a}")
        r_a.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # CHAPTER 10: CONCLUSION & FUTURE SCOPE
    # -------------------------------------------------------------
    add_chapter_heading("10", "Conclusion & Future Enhancements")
    
    p = doc.add_paragraph()
    p.add_run(
        "The University Internship Management & Tracking System successfully automates the critical workflows of academic placement "
        "divisions. By connecting students, placement officers, and industry supervisors within an intuitive, transparent interface, "
        "the portal eliminates manual paperwork delays, ensures authentic document verification, tracks supervisor evaluations, and "
        "coordinates recruitment drives seamlessly."
    )

    add_subheading("10.1 Future Scope")
    add_bullet(doc, "Automated Email & SMS Notifications", "Integrating Twilio / SendGrid for external email and SMS reminders for interview schedules.")
    add_bullet(doc, "External Industry Mentor Login", "A secure tokenized magic-link portal allowing external mentors to directly submit evaluations without account registration.")
    add_bullet(doc, "AI Resume Matcher", "Leveraging NLP to compute match scores between student resumes and published drive descriptions.")

    # Save document
    filename = "University_Internship_Management_System_Project_Report.docx"
    output_path = os.path.join(os.path.dirname(__file__), filename)
    doc.save(output_path)
    print(f"Document created successfully at: {output_path}")

if __name__ == '__main__':
    create_document()
